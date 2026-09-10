"""Build the representative 3-chapter fixture DOCX (V1 acceptance document).

Chapters follow spec section 41:
- Chapter A (Simple): headings + short paragraphs
- Chapter B (Dense): long paragraphs, multiple sections, lists
- Chapter C (Visual): images, definition, process explanation, concepts

The document also contains a table so the parser's unsupported-element
warning path is exercised.
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw


COURSE_TITLE = "Introduction to Artificial Intelligence"


def _make_image(path: Path) -> Path:
    """A 1600x900 raster image standing in for an embedded figure."""

    width, height = 1600, 900
    image = Image.new("RGB", (width, height), "#1E293B")
    draw = ImageDraw.Draw(image)
    for x in range(width):
        shade = int(30 + 80 * math.sin(x / 180))
        draw.line([(x, 0), (x, height)], fill=(30, 41 + shade // 3, 59 + shade // 2))
    draw.rounded_rectangle([240, 240, 1040, 480], radius=24, outline="#FF5A36", width=8)
    draw.ellipse([320, 300, 480, 460], fill="#0EA5E9", outline="#FFFFFF", width=6)
    draw.ellipse([520, 300, 680, 460], fill="#22C55E", outline="#FFFFFF", width=6)
    draw.ellipse([720, 300, 880, 460], fill="#F59E0B", outline="#FFFFFF", width=6)
    draw.line([(400, 300), (600, 460)], fill="#FFFFFF", width=6)
    draw.line([(600, 300), (800, 460)], fill="#FFFFFF", width=6)
    draw.line([(800, 300), (600, 460)], fill="#FFFFFF", width=6)
    draw.text((470, 120), "NEURAL NETWORK", fill="#FFFFFF", anchor="mm")
    image.save(path, "PNG")
    return path


def _long_paragraph() -> str:
    return (
        "Supervised learning is the branch of machine learning where a model is trained on "
        "labeled examples, meaning every training input has a known correct output. The "
        "learning process starts with a dataset that contains both features and labels. "
        "The model makes predictions on the training data and compares those predictions "
        "with the true labels using a loss function. The loss function measures how far "
        "the predictions are from the truth, and the training algorithm adjusts the model "
        "parameters to reduce that loss. This adjustment repeats over many iterations "
        "until the model performs well on the training set. Because the training process "
        "can memorize noise rather than real patterns, practitioners split their data "
        "into training and test sets. The test set contains examples the model has never "
        "seen, and it gives an honest estimate of how well the model generalizes to new "
        "data. This separation between fitting the data and evaluating on unseen data is "
        "one of the most important ideas in all of machine learning, and it applies "
        "equally to simple linear models and to very large neural networks."
    )


def build_document() -> Document:
    document = Document()
    document.core_properties.title = COURSE_TITLE
    document.core_properties.author = "Course Video Generator Fixture"

    document.add_heading("Introduction to Machine Learning", level=1)
    document.add_heading("What is Machine Learning?", level=2)
    document.add_paragraph(
        "Machine learning is a method that allows computers to learn patterns from data "
        "and use those patterns to make predictions without being explicitly programmed "
        "for every individual decision."
    )
    document.add_paragraph(
        "Instead of writing rules by hand, developers provide examples and the system "
        "finds the patterns on its own."
    )
    document.add_heading("Why It Matters", level=2)
    document.add_paragraph(
        "Machine learning powers recommendations, speech recognition, and medical "
        "diagnostics because it scales with data."
    )

    document.add_heading("Supervised Learning", level=1)
    document.add_heading("The Learning Process", level=2)
    document.add_paragraph(_long_paragraph())
    document.add_paragraph("The main ingredients of supervised learning are:")
    document.add_paragraph("Labeled training data", style="List Bullet")
    document.add_paragraph("A model with adjustable parameters", style="List Bullet")
    document.add_paragraph("A loss function that measures error", style="List Bullet")
    document.add_paragraph("An optimization algorithm", style="List Bullet")
    document.add_paragraph("A separate test set for honest evaluation", style="List Bullet")
    document.add_paragraph("A validation set for tuning choices", style="List Bullet")
    document.add_paragraph("Reproducible random seeds", style="List Bullet")
    document.add_heading("Common Algorithms", level=2)
    document.add_paragraph("Training a supervised model follows these steps:")
    document.add_paragraph("Split the data into training and test sets", style="List Number")
    document.add_paragraph("Feed a batch of examples through the model", style="List Number")
    document.add_paragraph("Compute the loss against the true labels", style="List Number")
    document.add_paragraph("Update the parameters to reduce the loss", style="List Number")
    document.add_paragraph(
        "The classic pipeline is Data -> Model -> Prediction, and each stage must be "
        "measured to find where errors enter the system."
    )

    document.add_heading("Neural Networks", level=1)
    document.add_heading("Anatomy of a Neural Network", level=2)
    document.add_paragraph(
        "A neural network is composed of layers of connected units called neurons."
    )
    image_path = _make_image(Path(__file__).parent / "_fixture_image.png")
    document.add_picture(str(image_path), width=Inches(6.0))
    document.paragraphs[-1].style = document.styles["Normal"]
    document.add_heading("What is a Neural Network?", level=2)
    document.add_paragraph(
        "A neural network is defined as a computational system of layered units that "
        "learns to transform inputs into outputs by adjusting connection weights during "
        "training."
    )
    document.add_heading("Old Approach vs New Approach", level=2)
    document.add_paragraph(
        "The old approach vs the new approach: hand-crafted rules for every case, "
        "while the new approach lets the network learn features automatically from "
        "raw data."
    )
    document.add_heading("Forward Propagation", level=2)
    document.add_paragraph(
        "During forward propagation, Input -> Hidden Layers -> Output, each layer "
        "transforms its input and passes the result to the next layer."
    )

    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Meaning"
    table.cell(1, 0).text = "Epoch"
    table.cell(1, 1).text = "One full pass over the training data"
    document.add_paragraph("Tables are not part of the V1 teaching output.")

    return document


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the fixture course DOCX")
    parser.add_argument("--out", default="input/course.docx")
    args = parser.parse_args()
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(str(out))
    print(f"Wrote fixture course to {out.resolve()}")


if __name__ == "__main__":
    main()
