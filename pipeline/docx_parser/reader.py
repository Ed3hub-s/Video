"""Order-preserving DOCX reader with unsupported-element warnings."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from pipeline.config import Config
from pipeline.content.schemas import Course

from .block_parser import ListKind, is_list_paragraph, parse_paragraph
from .chapter_parser import ChapterParser
from .image_extractor import detect_embedded_video, slugify


def _docx_text(paragraph) -> str:
    pieces: list[str] = []
    for child in paragraph._p.iter():
        if child.tag.split("}")[-1] == "t":
            pieces.append(child.text or "")
    return "".join(pieces).strip()


def _heading_level(paragraph) -> int | None:
    style_name = (paragraph.style.name or "") if paragraph.style else ""
    if style_name.lower().startswith("heading "):
        try:
            return int(style_name.rsplit(" ", 1)[1])
        except (ValueError, IndexError):
            return None
    outline = paragraph._p.pPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl"
    ) if paragraph._p.pPr is not None else None
    if outline is not None and outline.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"):
        try:
            return int(outline.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")) + 1
        except (ValueError, TypeError):
            return None
    return None


def _course_title(document: Document, input_path: Path) -> str:
    props = document.core_properties
    if props.title:
        return props.title.strip()
    for paragraph in document.paragraphs:
        style_name = (paragraph.style.name or "") if paragraph.style else ""
        if style_name.lower() == "title":
            text = _docx_text(paragraph)
            if text:
                return text
    return input_path.stem


def _check_document_warnings(document: Document, warnings: list[str]) -> None:
    body = document.element.body
    for child in body:
        if child.tag.endswith("}tbl"):
            warnings.append("Table found - tables are not supported in V1 and were skipped")
    if detect_embedded_video(document):
        warnings.append("Embedded video found - video is not supported in V1 and was skipped")
    for rel in document.part.rels.values():
        reltype = rel.reltype or ""
        if "footnotes" in reltype:
            warnings.append("Footnotes found - footnotes are not supported in V1 and were skipped")
        if "endnotes" in reltype:
            warnings.append("Endnotes found - endnotes are not supported in V1 and were skipped")
    xml = document.element.xml
    if "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ins" in xml:
        warnings.append("Tracked insertions found - tracked changes are not supported in V1")
    if "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}del" in xml:
        warnings.append("Tracked deletions found - tracked changes are not supported in V1")


def read_docx(input_path: Path, config: Config) -> Course:
    """Parse a .docx file into a Course, preserving document order."""

    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    document = Document(str(input_path))
    warnings: list[str] = []
    _check_document_warnings(document, warnings)
    title = _course_title(document, input_path)
    parser = ChapterParser(course_title=title, warnings=warnings)

    body = document.element.body
    paragraph_index = 0
    paragraphs = document.paragraphs
    paragraph_by_element = {p._p: p for p in paragraphs}
    image_counter = [0]

    for child in body:
        if child.tag.endswith("}p"):
            paragraph = paragraph_by_element.get(child)
            if paragraph is None:
                continue
            paragraph_index += 1
            level = _heading_level(paragraph)
            if level == 1:
                parser.add_heading1(_docx_text(paragraph))
                continue
            if level == 2:
                parser.add_heading2(_docx_text(paragraph))
                continue
            if level is not None and level > 2:
                parser.add_heading_other(_docx_text(paragraph), level)
                continue
            list_kind = is_list_paragraph(paragraph)
            if list_kind:
                parser.add_list_item(_docx_text(paragraph), list_kind)
                continue
            chapter = parser.chapters[-1] if parser.chapters else None
            chapter_dir = (
                config.paths.images / f"chapter-{chapter.chapter_number:02d}"
                if chapter
                else config.paths.images
            )
            blocks = parse_paragraph(
                paragraph,
                chapter_dir,
                image_counter,
                list_kind=ListKind.NONE,
            )
            for block in blocks:
                parser.add_block(block)

    return parser.finish()
